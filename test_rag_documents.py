"""rag_documents.py 的解析與目錄邏輯測試，不需要 data/ 也不會呼叫 OpenAI：
    python -m pytest test_rag_documents.py
"""

import docx as python_docx

import rag_documents as rd


def test_normalize_text_unifies_lookalike_characters():
    # U+F967 是相容字「不」，U+2F45 是康熙部首「⽅」，看起來一樣但編碼不同
    assert rd.normalize_text("不可以") == "不可以"
    assert rd.normalize_text("⽅法⼀") == "方法一"
    assert rd.normalize_text("ＴＯＥＩＣ６００") == "TOEIC600"


def test_normalize_text_keeps_chinese_punctuation_and_drops_private_use():
    assert rd.normalize_text("學費：（本籍生），一份20元") == "學費：（本籍生），一份20元"
    assert rd.normalize_text("造字測試") == "造字測試"


def test_rows_to_markdown_fills_down_first_column_merges():
    rows = [
        ["身分", "費用", "資電學院"],
        ["本籍生", "學費", "17,490"],
        [None, "雜費", "11,170"],
    ]
    lines = rd.rows_to_markdown(rows).split("\n")
    assert lines[0] == "| 身分 | 費用 | 資電學院 |"
    assert lines[3] == "| 本籍生 | 雜費 | 11,170 |"


def test_rows_to_markdown_joins_wrapped_cell_lines():
    rows = [["工、資電學院\n(含資管系)", "Student\nID"], ["a", "b"]]
    header = rd.rows_to_markdown(rows).split("\n")[0]
    assert "工、資電學院(含資管系)" in header
    assert "Student ID" in header


def test_rows_to_markdown_treats_frames_as_plain_text():
    # 每列只有一格有字的「表格」是版面外框，保留原本的換行
    assert rd.rows_to_markdown([["注意事項\n1. 請附正本"], ["", None], ["2. 審核需時7個工作天", ""]]) == (
        "注意事項\n1. 請附正本\n2. 審核需時7個工作天"
    )


def test_docx_merged_cells_are_not_repeated():
    document = python_docx.Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "資格考紀錄"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "演算法"
    table.cell(1, 1).text = "通過"
    table.cell(2, 1).text = "未通過"

    markdown = rd.rows_to_markdown(rd._docx_table_rows(table._tbl))
    assert markdown.count("資格考紀錄") == 1
    # 往下合併的第一欄會補回上一列的值，每一列才看得懂在講哪一科
    assert "| 演算法 | 未通過 |" in markdown


def test_parse_docx_keeps_paragraphs_and_tables_in_order(tmp_path):
    document = python_docx.Document()
    document.add_paragraph("工程五館門禁通行申請表")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "押金"
    table.cell(0, 1).text = "一百元"
    table.cell(1, 0).text = "期間"
    document.add_paragraph("歸還卡片即退還押金。")
    path = tmp_path / "門禁.docx"
    document.save(path)

    text, _ = rd.parse_docx(path)
    assert text.index("工程五館門禁通行申請表") < text.index("| 押金 | 一百元 |") < text.index("歸還卡片即退還押金。")


def test_continues_sentence_only_for_wrapped_lines():
    assert rd._continues_sentence("不含實驗室、電腦", "教室、語言教室")
    assert not rd._continues_sentence("特訂定本辦法。", "本辦法所稱之教室")
    assert not rd._continues_sentence("為可供一般課程使用之", "第二條、本辦法所稱")


def _doc(doc_id, title, version="", file_name=""):
    card = {"title": title, "version": version, "doc_type": "申請表單", "scope": "", "applies_to": "", "summary": "", "answers": []}
    return rd.CatalogDocument(doc_id=doc_id, file_name=file_name or f"{doc_id}.docx", text="", card=card)


def test_version_year_reads_roc_and_western_years():
    assert rd._version_year(_doc("D1", "申請表", version="114.06.10適用")) == 114
    assert rd._version_year(_doc("D1", "申請業務說明", version="2026/01/09 更新")) == 115
    assert rd._version_year(_doc("D1", "申請表", file_name="115專題確認表.docx")) == 115
    assert rd._version_year(_doc("D1", "申請表")) == 0


def test_older_versions_point_to_latest():
    catalog = [
        _doc("D01", "資訊工程學系112學年度「專題實驗」指導老師確認表"),
        _doc("D02", "資訊工程學系115學年度「專題實驗」指導老師確認表"),
        _doc("D03", "資訊工程學系113學年度「專題實驗」指導老師確認表"),
        _doc("D04", "學雜費收費標準", version="115學年度"),
    ]
    rd.mark_superseded_versions(catalog)
    assert [d.superseded_by for d in catalog] == ["D02", "", "D02", ""]


def test_versions_are_not_marked_when_years_are_unknown_or_tied():
    unknown = [_doc("D01", "文學院外文能力鑑定審核申請表"), _doc("D02", "文學院外文能力鑑定審核申請表")]
    tied = [_doc("D03", "離校同意書", version="114.10.15"), _doc("D04", "離校同意書", version="114.10.15")]
    rd.mark_superseded_versions(unknown + tied)
    assert all(d.superseded_by == "" for d in unknown + tied)


def test_markdown_snapshots_from_the_crawler_are_read_as_text(tmp_path, monkeypatch):
    monkeypatch.setattr(rd, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(rd, "CACHE_DIR", tmp_path / "cache")
    page = tmp_path / "data" / "學務處生活輔導組" / "校內獎學金一覽.md"
    page.parent.mkdir(parents=True)
    page.write_text("# 校內獎學金一覽\n\n| 名稱 | 金額 |\n| --- | --- |\n| 書卷獎 | 5,000 元 |\n", encoding="utf-8")

    assert rd.list_data_files() == [page]
    doc = rd.parse_file(page)
    # 爬蟲的檔案放在來源資料夾底下，檔名要帶資料夾，前端才連得到 /files/<來源>/<檔名>
    assert doc.file_name == "學務處生活輔導組/校內獎學金一覽.md"
    assert "| 書卷獎 | 5,000 元 |" in doc.text
