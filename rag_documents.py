"""校園法規 RAG 的文件處理：把 data/ 底下的 PDF / Word 檔轉成乾淨的文字，
再替每份文件產生一張「文件卡片」（標題、類型、適用範圍、版本、摘要、能回答
的問題），查詢時由 academic_agent.py 拿卡片目錄挑文件、讀全文回答。

產出都快取在 storage/rag/，以「檔案內容的 hash」當 key：檔案沒變就不會
重新解析、重新產生卡片；檔案刪掉，下次載入時自然就不在目錄裡了（舊版只看
最新修改時間，刪檔或放進舊檔都不會重建）。

踩過的坑（都是拿 data/ 裡的真實文件查出來的）：
- pdfplumber 解不了 Big5 字型，教室使用管理辦法.pdf 整份變亂碼；PyMuPDF 可以。
- 舊版「一頁有表格就只留表格」會把同一頁表格以外的文字整段丟掉，學雜費
  收費標準的學院欄位名稱、電子版成績單注意事項幾乎整頁都是這樣不見的。
  這裡表格跟表格外的文字都保留，依頁面上的位置排好。
- 很多公文用的是「CJK 相容字」或「康熙部首」（例如 ⽅ U+2F45 不是 方 U+65B9），
  看起來一樣但編碼不同，關鍵字比對會對不上，一律轉成一般的字。
- Word 的合併儲存格用 python-docx 的 row.cells 讀會重複出現很多次，
  改成直接讀 XML 裡實際存在的儲存格。
- 海報型的 PDF（在學證明申請說明.pdf）文字層順序是亂的，這種頁面才把
  頁面圖片跟文字層一起交給模型重排；一般頁面完全不經過 LLM。

直接執行 `python rag_documents.py` 會解析全部文件、補齊卡片，並印出目錄。
"""

import base64
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import docx as python_docx
import fitz  # PyMuPDF
from docx.oxml.ns import qn
from dotenv import load_dotenv
from openai import OpenAI

from logging_config import make_print_logger

print = make_print_logger(__name__)

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
CACHE_DIR = BASE_DIR / "storage" / "rag"
DOC_CONVERTED_CACHE_DIR = BASE_DIR / "doc_converted_cache"

# 解析或卡片的邏輯改了就把版本號加一，舊快取會自動失效。
PARSER_VERSION = "2"
CARD_VERSION = "2"

# .odt、.doc 用 LibreOffice 轉成 .docx 再讀，.md 是爬蟲把網頁內容存下來的（crawler.py）。
SUPPORTED_SUFFIXES = (".pdf", ".docx", ".doc", ".odt", ".md")

VISION_MODEL = os.getenv("RAG_VISION_MODEL", "gpt-5.4")
CARD_MODEL = os.getenv("RAG_CARD_MODEL", "gpt-5.4")

# 卡片只需要看懂文件在講什麼，太長的文件截斷前面這麼多字就夠了。
CARD_INPUT_MAX_CHARS = 24000

_client: Optional[OpenAI] = None


def _openai() -> OpenAI:
    global _client
    if _client is None:
        _client = OpenAI()
    return _client


# ============================================================
# 文字正規化
# ============================================================
def _needs_nfkc(cp: int) -> bool:
    return (
        0xF900 <= cp <= 0xFAFF  # CJK 相容表意字
        or 0x2F800 <= cp <= 0x2FA1F  # CJK 相容表意字補充
        or 0x2F00 <= cp <= 0x2FDF  # 康熙部首
        or 0x2E80 <= cp <= 0x2EFF  # CJK 部首補充
        or 0xFF10 <= cp <= 0xFF19  # 全形數字
        or 0xFF21 <= cp <= 0xFF3A  # 全形大寫英文
        or 0xFF41 <= cp <= 0xFF5A  # 全形小寫英文
    )


def normalize_text(text: str) -> str:
    """只轉「長得一樣、編碼不同」的字，中文標點（，：（）等）維持原樣。"""
    out = []
    for ch in text or "":
        cp = ord(ch)
        if _needs_nfkc(cp):
            out.append(unicodedata.normalize("NFKC", ch))
        elif ch == "\u3000":
            out.append(" ")
        elif 0xE000 <= cp <= 0xF8FF or ch in "\u200b\ufeff" or (cp < 32 and ch not in "\n\t"):
            continue  # 私用區（造字）跟控制字元印出來都是亂碼，直接丟掉
        else:
            out.append(ch)
    text = "".join(out)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


_CJK_CHAR_RE = re.compile(r"[\u3400-\u9fff\uff08\uff09\u3001-\u3011，：；]")


def _join_cell_lines(text: str) -> str:
    """儲存格裡的換行只是排版折行：中文後面接中文或括號直接接起來，其他情況補空格。"""
    lines = [line.strip() for line in (text or "").split("\n") if line.strip()]
    if not lines:
        return ""
    merged = lines[0]
    for line in lines[1:]:
        if _CJK_CHAR_RE.match(merged[-1]) and (_CJK_CHAR_RE.match(line[0]) or line[0] == "("):
            merged += line
        else:
            merged += " " + line
    return merged.replace("|", "｜")


def rows_to_markdown(rows: list[list[Optional[str]]]) -> str:
    """表格列轉 Markdown。None 代表合併儲存格：第一欄往下合併的（例如
    「本籍生」跨好幾列）把上一列的值補回來，每一列才看得懂自己在講誰；
    其他欄位的合併就留空，避免把橫跨整列的標題重複很多次。
    """
    if all(sum(1 for c in row if c and c.strip()) <= 1 for row in rows):
        # 每一列最多一格有字的「表格」多半是版面外框（注意事項、公告框），
        # 當一般文字、保留原本的換行比較好讀。
        return "\n".join(line.strip() for row in rows for c in row if c for line in c.split("\n") if line.strip())

    cleaned: list[list[str]] = []
    previous_first = ""
    for row in rows:
        cells = [_join_cell_lines(c) if c is not None else None for c in row]
        if cells and cells[0] is None:
            cells[0] = previous_first
        elif cells:
            previous_first = cells[0]
        cells = [c or "" for c in cells]
        if any(cells):
            cleaned.append(cells)

    if not cleaned:
        return ""
    width = max(len(r) for r in cleaned)
    cleaned = [r + [""] * (width - len(r)) for r in cleaned]
    lines = ["| " + " | ".join(cleaned[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in cleaned[1:]]
    return "\n".join(lines)


# ============================================================
# PDF
# ============================================================
def _image_coverage(page) -> float:
    area = page.rect.width * page.rect.height or 1.0
    covered = 0.0
    for info in page.get_image_info():
        r = fitz.Rect(info["bbox"]) & page.rect
        if not r.is_empty:
            covered += r.width * r.height
    return min(covered / area, 1.0)


def _looks_garbled(text: str, file_name: str) -> bool:
    """中文檔名的文件卻幾乎抽不到中文字，通常是字型編碼解錯了。"""
    body = re.sub(r"\s", "", text)
    if len(body) < 30 or not re.search(r"[\u4e00-\u9fff]", file_name):
        return False
    cjk = len(re.findall(r"[\u4e00-\u9fff]", body))
    return cjk / len(body) < 0.05


def _page_needs_vision(page, text: str, file_name: str) -> bool:
    return _image_coverage(page) >= 0.5 or _looks_garbled(text, file_name)


_PARAGRAPH_START_RE = re.compile(
    r"^(第[一二三四五六七八九十百零〇0-9]+[條章節項款]|[一二三四五六七八九十壹貳參肆伍陸柒捌玖拾]+[、.]"
    r"|[（(][一二三四五六七八九十0-9]+[)）]|[0-9]+[.、)）]|[■□◆●○※★•‧▪\-－]|附[表件錄]|註)"
)


def _continues_sentence(previous: str, following: str) -> bool:
    """PDF 常把同一句話折行後切成兩個文字區塊（「…不含實驗室、電腦」「教室、
    語言教室…」），這種情況不該當成新段落。
    """
    return (
        bool(previous) and bool(following)
        and re.match(r"[㐀-鿿，、（(「]", previous[-1]) is not None
        and re.match(r"[㐀-鿿）)」，、。]", following[0]) is not None
        and not _PARAGRAPH_START_RE.match(following)
    )


def _pdf_page_markdown(page) -> str:
    """表格轉 Markdown，表格外的文字區塊照原樣保留，兩者依頁面位置排序。"""
    try:
        tables = page.find_tables().tables
    except Exception as e:
        print(f"[PDF] 第 {page.number + 1} 頁找表格失敗，只用純文字：{e}")
        tables = []

    items: list[tuple[float, float, str, bool]] = []  # (y, x, 內容, 是不是表格)
    table_rects = []
    for table in tables:
        markdown = rows_to_markdown(table.extract())
        if markdown:
            items.append((table.bbox[1], table.bbox[0], markdown, True))
            table_rects.append(fitz.Rect(table.bbox))

    for x0, y0, x1, y1, text, _, block_type in page.get_text("blocks", sort=True):
        if block_type != 0 or not text.strip():
            continue
        rect = fitz.Rect(x0, y0, x1, y1)
        block_area = rect.width * rect.height or 1.0
        if any((rect & t).width * (rect & t).height / block_area > 0.6 for t in table_rects if rect.intersects(t)):
            continue  # 表格裡的文字已經在 Markdown 表格裡了
        items.append((y0, x0, text.strip(), False))

    items.sort(key=lambda item: (round(item[0]), item[1]))
    output = ""
    previous_is_table = True
    for _, _, text, is_table in items:
        if not output:
            output = text
        elif not is_table and not previous_is_table and _continues_sentence(output, text):
            output += "\n" + text
        else:
            output += "\n\n" + text
        previous_is_table = is_table
    return output


VISION_PROMPT = """這是一份中央大學校務文件的其中一頁（可能是海報或版面很複雜的頁面）。
請把圖片裡的內容完整轉錄成閱讀順序正確的 Markdown：

1. 以圖片為準。下面附的「文字層」是從 PDF 抽出來的，順序可能是亂的，也可能有重複字或
   缺字，只拿來確認圖片裡不好辨識的字。
2. 不要摘要、不要補充頁面上沒有的內容、不要加開場白或結尾說明。
3. 表格用 Markdown 表格，分區或步驟用標題或清單。

文字層：
{text_layer}"""


def _vision_transcribe(page, text_layer: str) -> str:
    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2))
    image_b64 = base64.b64encode(pixmap.tobytes("png")).decode("ascii")
    response = _openai().chat.completions.create(
        model=VISION_MODEL,
        reasoning_effort="low",
        messages=[{
            "role": "user",
            "content": [
                {"type": "text", "text": VISION_PROMPT.format(text_layer=text_layer or "（沒有文字層）")},
                {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{image_b64}", "detail": "high"}},
            ],
        }],
    )
    result = (response.choices[0].message.content or "").strip()
    result = re.sub(r"^```(?:markdown)?\s*|\s*```$", "", result)
    return result


def parse_pdf(path: Path) -> tuple[str, int]:
    pages = []
    with fitz.open(path) as pdf:
        for page in pdf:
            text = _pdf_page_markdown(page)
            if _page_needs_vision(page, text, path.name):
                print(f"[PDF] {path.name} 第 {page.number + 1} 頁是海報/圖片或文字層亂碼，改用圖片輔助轉錄")
                try:
                    text = _vision_transcribe(page, page.get_text("text", sort=True))
                except Exception as e:
                    print(f"[PDF] {path.name} 第 {page.number + 1} 頁圖片轉錄失敗，保留原本的文字層：{e}")
            pages.append(f"〔第 {page.number + 1} 頁〕\n{normalize_text(text)}")
        page_count = pdf.page_count
    return "\n\n".join(pages), page_count


# ============================================================
# Word（.docx，以及透過 LibreOffice 轉成 .docx 的 .doc、.odt）
# ============================================================
def _xml_text(element) -> str:
    parts = []
    for node in element.iter(qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")):
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:tab"):
            parts.append(" ")
        else:
            parts.append("\n")
    return "".join(parts)


def _docx_cell_text(tc) -> str:
    lines = []
    for child in tc.iterchildren():
        if child.tag == qn("w:p"):
            lines.append(_xml_text(child))
        elif child.tag == qn("w:tbl"):  # 儲存格裡的巢狀表格，攤平成一列一行
            for tr in child.iterchildren(qn("w:tr")):
                lines.append(" ".join(_docx_cell_text(c).replace("\n", " ") for c in tr.iterchildren(qn("w:tc"))))
    return "\n".join(line for line in lines if line.strip())


def _docx_table_rows(tbl) -> list[list[Optional[str]]]:
    rows = []
    for tr in tbl.iterchildren(qn("w:tr")):
        row: list[Optional[str]] = []
        for tc in tr.iterchildren(qn("w:tc")):
            tc_pr = tc.find(qn("w:tcPr"))
            span, continues_above = 1, False
            if tc_pr is not None:
                grid_span = tc_pr.find(qn("w:gridSpan"))
                if grid_span is not None:
                    span = int(grid_span.get(qn("w:val"), "1"))
                v_merge = tc_pr.find(qn("w:vMerge"))
                continues_above = v_merge is not None and v_merge.get(qn("w:val")) != "restart"
            # 往下合併的延續格用 None，跟 PDF 表格一樣交給 rows_to_markdown 處理；
            # 橫向合併只輸出一次內容，其餘位置補空格讓欄位對齊。
            row.append(None if continues_above else _docx_cell_text(tc))
            row.extend([""] * (span - 1))
        rows.append(row)
    return rows


def _docx_blocks(parent) -> list[str]:
    blocks = []
    for child in parent.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = _xml_text(child).strip()
            if text:
                blocks.append(text)
        elif tag == "tbl":
            markdown = rows_to_markdown(_docx_table_rows(child))
            if markdown:
                blocks.append("\n" + markdown + "\n")
        elif tag == "sdt":  # 內容控制項（表單欄位），裡面還是一般段落/表格
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                blocks.extend(_docx_blocks(content))
    return blocks


def parse_docx(path: Path) -> tuple[str, int]:
    document = python_docx.Document(str(path))
    return normalize_text("\n".join(_docx_blocks(document.element.body))), 0


def _find_soffice() -> Optional[str]:
    configured = os.getenv("LIBREOFFICE_PATH")
    if configured and os.path.exists(configured):
        return configured
    for candidate in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ):
        if os.path.exists(candidate):
            return candidate
    return shutil.which("soffice")


def convert_to_docx(path: Path, file_hash: str) -> Optional[Path]:
    """舊版 .doc、OpenDocument 的 .odt 用 LibreOffice 無頭轉檔成 .docx，轉不了回傳 None。
    結果以檔案內容的 hash 快取：爬蟲會在不同資料夾放同名檔案，用檔名當 key 會互相蓋掉。
    """
    DOC_CONVERTED_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cached = DOC_CONVERTED_CACHE_DIR / f"{file_hash}.docx"
    if cached.exists():
        return cached

    soffice = _find_soffice()
    if not soffice:
        print(
            f"[{path.name}] 找不到 LibreOffice，沒辦法轉 {path.suffix}。請安裝 "
            "https://www.libreoffice.org/download/download/ 或設定 LIBREOFFICE_PATH。"
        )
        return None

    with tempfile.TemporaryDirectory() as work_dir:
        # 複製成 ASCII 檔名再轉：LibreOffice 遇到某些中文、全形符號的路徑會轉檔失敗
        source = Path(work_dir) / f"input{path.suffix.lower()}"
        shutil.copyfile(path, source)
        try:
            result = subprocess.run(
                [soffice, "--headless", "--convert-to", "docx", "--outdir", work_dir, str(source)],
                capture_output=True, text=True, timeout=120,
            )
        except subprocess.TimeoutExpired:
            print(f"[{path.name}] LibreOffice 轉檔超過 120 秒，略過這個檔案。")
            return None
        converted = Path(work_dir) / "input.docx"
        if result.returncode != 0 or not converted.exists():
            print(f"[{path.name}] LibreOffice 轉檔失敗：{result.stderr.strip() or result.stdout.strip()}")
            return None
        shutil.move(str(converted), str(cached))
    return cached


# ============================================================
# 解析 + 快取
# ============================================================
@dataclass
class ParsedDocument:
    file_name: str
    text: str
    page_count: int
    file_hash: str

    @property
    def content_key(self) -> str:
        """內容一樣的檔案（例如同一份表單存成兩個檔名）拿到同一個 key。"""
        return hashlib.sha256(re.sub(r"\s", "", self.text).encode("utf-8")).hexdigest()


def _file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_json(path: Path) -> Optional[dict]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def _write_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(path)


def parse_file(path: Path) -> Optional[ParsedDocument]:
    # 前端會把來源做成 /files/<檔名> 的連結（main.py 把 data/ 掛在 /files），
    # 所以這裡用 data/ 底下的相對路徑當檔名。
    file_name = path.relative_to(DATA_DIR).as_posix()
    file_hash = _file_sha256(path)
    cache_path = CACHE_DIR / "parsed" / f"{file_hash}.json"
    cached = _read_json(cache_path)
    if cached and cached.get("parser_version") == PARSER_VERSION:
        return ParsedDocument(file_name, cached["text"], cached["page_count"], file_hash)

    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            text, page_count = parse_pdf(path)
        elif suffix == ".docx":
            text, page_count = parse_docx(path)
        elif suffix in (".doc", ".odt"):
            converted = convert_to_docx(path, file_hash)
            if converted is None:
                return None
            text, page_count = parse_docx(converted)
        elif suffix == ".md":
            text, page_count = normalize_text(path.read_text(encoding="utf-8")), 0
        else:
            return None
    except Exception as e:
        print(f"[文件解析] {path.name} 解析失敗，略過：{e}")
        return None

    if not re.sub(r"〔第 \d+ 頁〕|\s", "", text):
        print(f"[文件解析] {path.name} 抽不到任何文字，略過。")
        return None

    _write_json(cache_path, {"parser_version": PARSER_VERSION, "file_name": file_name, "text": text, "page_count": page_count})
    return ParsedDocument(file_name, text, page_count, file_hash)


def list_data_files() -> list[Path]:
    if not DATA_DIR.exists():
        return []
    return sorted(
        p for p in DATA_DIR.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES and not p.name.startswith(("~$", "."))
    )


def data_signature() -> tuple:
    """用來判斷 data/ 有沒有變（新增、刪除、修改都算），比對很便宜。"""
    return tuple((str(p.relative_to(DATA_DIR)), p.stat().st_size, p.stat().st_mtime_ns) for p in list_data_files())


# ============================================================
# 文件卡片
# ============================================================
DOC_TYPES = ["法規辦法", "申請表單", "申請說明", "常見問答", "其他"]

CARD_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "doc_type": {"type": "string", "enum": DOC_TYPES},
        "issuer": {"type": "string"},
        "scope": {"type": "string"},
        "applies_to": {"type": "string"},
        "version": {"type": "string"},
        "summary": {"type": "string"},
        "answers": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["title", "doc_type", "issuer", "scope", "applies_to", "version", "summary", "answers"],
    "additionalProperties": False,
}

CARD_PROMPT = """你要替中央大學校園法規問答系統整理一份文件的「目錄卡片」。之後系統會只看卡片
決定要不要打開這份文件回答使用者，所以卡片要讓人一看就知道這份文件「是什麼、
適用於誰、哪個版本、能回答哪些問題」。以文件內容為準，檔名可以當參考（例如檔名
寫了適用學年度），看不出來的欄位填空字串，不要猜。

欄位說明：
- title：文件的正式名稱，以文件內文的標題為準（檔名可能不準），有年度或適用範圍要保留。
- doc_type：法規辦法（辦法、要點、細則、準則、規定、收費標準）、申請表單（要填寫的表格）、
  申請說明（申請流程、注意事項、公告）、常見問答、其他。
- issuer：訂定或承辦的單位，例如「教務處」「資訊工程學系」「工學院」。
- scope：適用範圍，例如「全校」「工學院」「資訊工程學系」「台灣聯合大學系統」。
- applies_to：適用或需要用到這份文件的人，例如「大學部學生」「碩士班新生」「博士班學生」
  「授課教師」「外籍生」「社會人士」，可以寫多個。
- version：版本或生效時間，例如「115學年度」「114學年度後入學新生適用」「114.06.10 起適用」
  「112學年度」；文件沒寫就空字串，不要自己推測。
- summary：兩到三句話說明文件內容與用途（表單要說明用來申請什麼、要誰簽核、繳交給誰）。
- answers：這份文件能回答的具體問題，用使用者會問的口吻寫，5 到 12 題；
  常見問答類的文件要把每一題都列進來（可以精簡措辭）。問題裡要帶出具體的主題詞
  （例如「電機系學生修資管系的演算法可以抵本系必修嗎」而不是「可以抵免嗎」）。

檔名：{file_name}

文件內容：
{text}"""


def build_card(doc: ParsedDocument) -> dict:
    text = doc.text
    if len(text) > CARD_INPUT_MAX_CHARS:
        text = text[:CARD_INPUT_MAX_CHARS] + "\n（以下省略）"
    response = _openai().chat.completions.create(
        model=CARD_MODEL,
        reasoning_effort="low",
        messages=[{"role": "user", "content": CARD_PROMPT.format(file_name=doc.file_name, text=text)}],
        response_format={"type": "json_schema", "json_schema": {"name": "document_card", "schema": CARD_SCHEMA, "strict": True}},
    )
    return json.loads(response.choices[0].message.content)


def _cached_card(doc: ParsedDocument) -> Optional[dict]:
    cached = _read_json(CACHE_DIR / "cards" / f"{doc.content_key}.json")
    if cached and cached.get("card_version") == CARD_VERSION:
        return cached["card"]
    return None


def _build_and_cache_card(doc: ParsedDocument) -> Optional[dict]:
    try:
        card = build_card(doc)
    except Exception as e:
        print(f"[文件卡片] {doc.file_name} 產生卡片失敗：{e}")
        return None
    _write_json(CACHE_DIR / "cards" / f"{doc.content_key}.json", {"card_version": CARD_VERSION, "file_name": doc.file_name, "card": card})
    return card


def _fallback_card(doc: ParsedDocument) -> dict:
    """卡片產生失敗時（例如沒網路）的最低限度卡片，至少讓檔名能被挑到。"""
    title = Path(doc.file_name).stem
    return {
        "title": title, "doc_type": "其他", "issuer": "", "scope": "", "applies_to": "", "version": "",
        "summary": doc.text[:120].replace("\n", " "), "answers": [],
    }


@dataclass
class CatalogDocument:
    doc_id: str
    file_name: str
    text: str
    card: dict
    duplicates: list[str] = field(default_factory=list)
    # 同一份文件有更新的年度版本時，填最新版的 doc_id（例如 112～115 學年度的專題確認表）。
    superseded_by: str = ""


def _series_key(title: str) -> str:
    """去掉年度、學期、數字跟標點後的標題，一樣的就當成同一份文件的不同版本。"""
    title = re.sub(r"\d+\s*學年度?(\s*第\s*\d+\s*學期)?", "", title)
    return re.sub(r"[\s\d_\-【】\[\]()（）「」『』]", "", title)


def _version_year(doc: "CatalogDocument") -> int:
    """從版本、標題、檔名裡找民國年（112、114.06.10）或西元年（2026），取最大的，找不到回傳 0。"""
    text = f"{doc.card['version']} {doc.card['title']} {doc.file_name}"
    years = [int(y) for y in re.findall(r"(?<!\d)(1[0-4]\d)(?!\d)", text)]
    years += [int(y) - 1911 for y in re.findall(r"(?<!\d)(20\d\d)(?!\d)", text)]
    return max(years, default=0)


def mark_superseded_versions(catalog: list["CatalogDocument"]) -> None:
    """同系列文件裡年度最新的那份以外，都標上 superseded_by。年度一樣或看不出年度就不標，
    交給挑文件的模型自己看內容判斷。
    """
    series: dict[str, list[CatalogDocument]] = {}
    for doc in catalog:
        series.setdefault(_series_key(doc.card["title"]), []).append(doc)
    for docs in series.values():
        if len(docs) < 2:
            continue
        latest = max(docs, key=_version_year)
        latest_year = _version_year(latest)
        if latest_year == 0 or sum(1 for d in docs if _version_year(d) == latest_year) > 1:
            continue
        for doc in docs:
            if doc is not latest:
                doc.superseded_by = latest.doc_id


def load_catalog(max_workers: int = 8) -> list[CatalogDocument]:
    """解析 data/ 全部文件（有快取）、合併內容完全相同的檔案、補齊文件卡片。"""
    parsed = [doc for doc in (parse_file(p) for p in list_data_files()) if doc is not None]

    unique: dict[str, ParsedDocument] = {}
    duplicates: dict[str, list[str]] = {}
    for doc in parsed:
        key = doc.content_key
        if key in unique:
            duplicates.setdefault(key, []).append(doc.file_name)
            print(f"[文件目錄] {doc.file_name} 跟 {unique[key].file_name} 內容相同，只保留一份")
        else:
            unique[key] = doc

    docs = list(unique.values())
    cards = {doc.content_key: _cached_card(doc) for doc in docs}
    missing = [doc for doc in docs if cards[doc.content_key] is None]
    if missing:
        print(f"[文件卡片] 有 {len(missing)} 份文件還沒有卡片，開始產生...")
        with ThreadPoolExecutor(max_workers=max_workers) as pool:
            for doc, card in zip(missing, pool.map(_build_and_cache_card, missing)):
                cards[doc.content_key] = card

    catalog = []
    for i, doc in enumerate(docs, 1):
        catalog.append(CatalogDocument(
            doc_id=f"D{i:02d}",
            file_name=doc.file_name,
            text=doc.text,
            card=cards[doc.content_key] or _fallback_card(doc),
            duplicates=duplicates.get(doc.content_key, []),
        ))
    mark_superseded_versions(catalog)
    return catalog


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    for entry in load_catalog():
        card = entry.card
        old = f"｜舊版→{entry.superseded_by}" if entry.superseded_by else ""
        print(f"{entry.doc_id} {card['title']}｜{card['doc_type']}｜{card['scope']}｜{card['version']}｜{entry.file_name}{old}")
