"""rag_documents.py 的解析與目錄邏輯測試，不需要 data/ 也不會呼叫 OpenAI：
    python -m pytest tests/test_rag_documents.py
"""

import docx as python_docx

from backend.rag import rag_documents as rd


def test_normalize_text_unifies_lookalike_characters():
    # U+F967 是相容字「不」，U+2F45 是康熙部首「⽅」，看起來一樣但編碼不同
    assert rd.normalize_text("不可以") == "不可以"
    assert rd.normalize_text("⽅法⼀") == "方法一"
    assert rd.normalize_text("ＴＯＥＩＣ６００") == "TOEIC600"


def test_normalize_text_keeps_chinese_punctuation_and_drops_private_use():
    assert rd.normalize_text("學費：（本籍生），一份20元") == "學費：（本籍生），一份20元"
    assert rd.normalize_text("造字測試") == "造字測試"


def test_rows_to_markdown_fills_down_first_column_merges():
    rows = [
        ["身分", "費用", "資電學院"],
        ["本籍生", "學費", "17,490"],
        [None, "雜費", "11,170"],
    ]
    lines = rd.rows_to_markdown(rows).split("\n")
    assert lines[0] == "| 身分 | 費用 | 資電學院 |"
    assert lines[3] == "| 本籍生 | 雜費 | 11,170 |"


def test_rows_to_markdown_fills_down_merges_in_every_column():
    # 圖書館借閱規則：博士班那格「30」往下合併到碩士班，不補的話碩士班看起來沒有借期
    rows = [
        ["身分別", "借閱總筆數", "借閱日數", "備註"],
        ["博士班學生", "100", "30", "逾期每日每冊罰五元，逾期三十日以上或累計三百元以上，暫停各項借用權利"],
        ["碩士班學生", "100", None, None],
        ["教職員工配偶", "20", "60", ""],
    ]
    lines = rd.rows_to_markdown(rows).split("\n")
    assert lines[3] == "| 碩士班學生 | 100 | 30（與上列合併） | （同上） |"  # 太長的內容不重複，寫「同上」
    assert lines[4] == "| 教職員工配偶 | 20 | 60 |  |"


def test_merged_subtotals_are_marked_instead_of_repeated():
    # 共同必修：「通識課程 14 學分」跨核心必修、選修科目兩列，是兩列合起來 14，不能看起來像各 14
    rows = [["科目", "類別", "學分數"], ["通識課程", "核心必修", "14"], [None, "選修科目", None], ["合計", "", "25"]]
    assert "| 通識課程 | 選修科目 | 14（與上列合併） |" in rd.rows_to_markdown(rows)


def test_pdf_tables_tell_downward_merges_from_sideways_ones():
    from types import SimpleNamespace

    def row(top, bottom, *cells):
        return SimpleNamespace(bbox=(0, top, 300, bottom), cells=list(cells))

    # 第二列：「借閱日數」那格從第一列往下延伸（往下合併），「說明」被左邊的格子橫向蓋住
    table = SimpleNamespace(
        extract=lambda: [["博士班", "30", "限館內"], ["碩士班", None, None]],
        rows=[row(0, 10, (0, 0, 100, 10), (100, 0, 200, 20), (200, 0, 300, 10)),
              row(10, 20, (0, 10, 100, 20), None, None)],
    )
    assert rd._pdf_table_rows(table) == [["博士班", "30", "限館內"], ["碩士班", None, ""]]


def test_reparsing_a_pdf_reuses_earlier_image_transcriptions(tmp_path, monkeypatch):
    import fitz

    path = tmp_path / "海報.pdf"
    with fitz.open() as pdf:
        pdf.new_page().insert_text((72, 72), "poster")
        pdf.save(path)

    def no_vision(page, text_layer):
        raise AssertionError("不該重新送去轉錄")

    monkeypatch.setattr(rd, "_page_needs_vision", lambda page, text, name: True)
    monkeypatch.setattr(rd, "_vision_transcribe", no_vision)
    text, _ = rd.parse_pdf(path, rd.split_pages("〔第 1 頁〕\n上次轉錄的海報內容"))
    assert text == "〔第 1 頁〕\n上次轉錄的海報內容"


def test_split_pages_reads_back_parsed_pdf_text():
    assert rd.split_pages("〔第 1 頁〕\n第一頁\n\n〔第 2 頁〕\n第二頁\n第二行") == {1: "第一頁", 2: "第二頁\n第二行"}


def test_rows_to_markdown_joins_wrapped_cell_lines():
    rows = [["工、資電學院\n(含資管系)", "Student\nID"], ["a", "b"]]
    header = rd.rows_to_markdown(rows).split("\n")[0]
    assert "工、資電學院(含資管系)" in header
    assert "Student ID" in header


def test_rows_to_markdown_treats_frames_as_plain_text():
    # 每列只有一格有字的「表格」是版面外框，保留原本的換行
    assert rd.rows_to_markdown([["注意事項\n1. 請附正本"], ["", None], ["2. 審核需時7個工作天", ""]]) == (
        "注意事項\n1. 請附正本\n2. 審核需時7個工作天"
    )


def test_docx_merged_cells_are_not_repeated():
    document = python_docx.Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "資格考紀錄"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "演算法"
    table.cell(1, 1).text = "通過"
    table.cell(2, 1).text = "未通過"

    markdown = rd.rows_to_markdown(rd._docx_table_rows(table._tbl))
    assert markdown.count("資格考紀錄") == 1
    # 往下合併的會補回上一列的值，每一列才看得懂在講哪一科
    assert "| 演算法 | 未通過 |" in markdown


def test_docx_downward_merges_fill_every_column():
    document = python_docx.Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "博士班"
    table.cell(1, 0).text = "碩士班"
    table.cell(2, 0).text = "大學部"
    table.cell(0, 1).merge(table.cell(1, 1)).text = "30 天"
    table.cell(2, 1).text = "14 天"
    markdown = rd.rows_to_markdown(rd._docx_table_rows(table._tbl))
    assert "| 碩士班 | 30 天（與上列合併） |" in markdown and "| 大學部 | 14 天 |" in markdown


def test_parse_docx_keeps_paragraphs_and_tables_in_order(tmp_path):
    document = python_docx.Document()
    document.add_paragraph("工程五館門禁通行申請表")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "押金"
    table.cell(0, 1).text = "一百元"
    table.cell(1, 0).text = "期間"
    document.add_paragraph("歸還卡片即退還押金。")
    path = tmp_path / "門禁.docx"
    document.save(path)

    text, _ = rd.parse_docx(path)
    assert text.index("工程五館門禁通行申請表") < text.index("| 押金 | 一百元 |") < text.index("歸還卡片即退還押金。")


def test_continues_sentence_only_for_wrapped_lines():
    assert rd._continues_sentence("不含實驗室、電腦", "教室、語言教室")
    assert not rd._continues_sentence("特訂定本辦法。", "本辦法所稱之教室")
    assert not rd._continues_sentence("為可供一般課程使用之", "第二條、本辦法所稱")


def _doc(doc_id, title, version="", file_name=""):
    card = {"title": title, "version": version, "doc_type": "申請表單", "scope": "", "applies_to": "", "summary": "", "answers": []}
    return rd.CatalogDocument(doc_id=doc_id, file_name=file_name or f"{doc_id}.docx", text="", card=card)


def test_version_year_reads_roc_and_western_years():
    assert rd._version_year(_doc("D1", "申請表", version="114.06.10適用")) == 114
    assert rd._version_year(_doc("D1", "申請業務說明", version="2026/01/09 更新")) == 115
    assert rd._version_year(_doc("D1", "申請表", file_name="115專題確認表.docx")) == 115
    assert rd._version_year(_doc("D1", "申請表")) == 0


def test_older_versions_point_to_latest():
    catalog = [
        _doc("D01", "資訊工程學系112學年度「專題實驗」指導老師確認表"),
        _doc("D02", "資訊工程學系115學年度「專題實驗」指導老師確認表"),
        _doc("D03", "資訊工程學系113學年度「專題實驗」指導老師確認表"),
        _doc("D04", "學雜費收費標準", version="115學年度"),
    ]
    rd.mark_superseded_versions(catalog)
    assert [d.superseded_by for d in catalog] == ["D02", "", "D02", ""]


def test_versions_are_not_marked_when_years_are_unknown_or_tied():
    unknown = [_doc("D01", "文學院外文能力鑑定審核申請表"), _doc("D02", "文學院外文能力鑑定審核申請表")]
    tied = [_doc("D03", "離校同意書", version="114.10.15"), _doc("D04", "離校同意書", version="114.10.15")]
    rd.mark_superseded_versions(unknown + tied)
    assert all(d.superseded_by == "" for d in unknown + tied)


def test_markdown_snapshots_from_the_crawler_are_read_as_text(tmp_path, monkeypatch):
    monkeypatch.setattr(rd, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(rd, "CACHE_DIR", tmp_path / "cache")
    page = tmp_path / "data" / "學務處生活輔導組" / "校內獎學金一覽.md"
    page.parent.mkdir(parents=True)
    page.write_text("# 校內獎學金一覽\n\n| 名稱 | 金額 |\n| --- | --- |\n| 書卷獎 | 5,000 元 |\n", encoding="utf-8")

    assert rd.list_data_files() == [page]
    doc = rd.parse_file(page)
    # 爬蟲的檔案放在來源資料夾底下，檔名要帶資料夾，前端才連得到 /files/<來源>/<檔名>
    assert doc.file_name == "學務處生活輔導組/校內獎學金一覽.md"
    assert "| 書卷獎 | 5,000 元 |" in doc.text


def _catalog_doc(doc_id, title, scope="全校", folder="教務處"):
    card = {"title": title, "doc_type": "法規辦法", "issuer": "", "scope": scope, "applies_to": "", "version": "",
            "summary": "摘要", "answers": ["問題？"]}
    return rd.CatalogDocument(doc_id, f"{folder}/{title}.pdf", "全文", card)


def test_versions_with_different_scopes_are_separate_series():
    # 兩個系的「碩士班修業規定」標題一樣，不能把物理系的舊版指到化學系的新版
    physics = _catalog_doc("D01", "112學年度碩士班修業規定", "物理學系", "物理學系")
    chemistry = _catalog_doc("D02", "114學年度碩士班修業規定", "化學學系", "化學學系")
    # 同一份全校性辦法放在兩個單位的網站上、一新一舊，還是要標出舊版
    old = _catalog_doc("D03", "112學年度學分抵免辦法", folder="教務處課務組")
    new = _catalog_doc("D04", "114學年度學分抵免辦法", folder="教務處")
    rd.mark_superseded_versions([physics, chemistry, old, new])
    assert (physics.superseded_by, chemistry.superseded_by) == ("", "")
    assert old.superseded_by == "D04"


def test_card_embeddings_are_cached_by_card_content(tmp_path, monkeypatch):
    import numpy as np

    monkeypatch.setattr(rd, "CACHE_DIR", tmp_path)
    embedded = []

    def embed(texts):
        embedded.extend(texts)
        return np.ones((len(texts), 4), dtype=np.float32) / 2

    monkeypatch.setattr(rd, "embed_texts", embed)
    docs = [_catalog_doc("D01", "學則"), _catalog_doc("D02", "選課辦法")]
    rd.attach_embeddings(docs)
    assert len(embedded) == 2 and all(d.embedding is not None for d in docs)

    docs[1].card = {**docs[1].card, "summary": "改過的摘要"}
    again = [_catalog_doc("D01", "學則"), docs[1]]
    rd.attach_embeddings(again)
    assert len(embedded) == 3 and "改過的摘要" in embedded[-1]  # 只重算卡片有變的那份
    assert again[0].embedding is not None


def test_failed_embeddings_leave_documents_without_vectors(tmp_path, monkeypatch):
    def broken(texts):
        raise RuntimeError("沒有網路")

    monkeypatch.setattr(rd, "CACHE_DIR", tmp_path)
    monkeypatch.setattr(rd, "embed_texts", broken)
    docs = [_catalog_doc("D01", "學則")]
    rd.attach_embeddings(docs)
    assert docs[0].embedding is None


def test_answer_embeddings_are_cached_per_question(tmp_path, monkeypatch):
    import numpy as np

    monkeypatch.setattr(rd, "CACHE_DIR", tmp_path)
    embedded = []

    def embed(texts, dimensions=4):
        embedded.extend(texts)
        return np.ones((len(texts), dimensions), dtype=np.float32) / dimensions ** 0.5

    monkeypatch.setattr(rd, "embed_texts", embed)
    docs = [_catalog_doc("D01", "學則"), _catalog_doc("D02", "選課辦法")]
    docs[1].card = {**docs[1].card, "answers": ["可以加選幾學分？", "停修有什麼限制？"]}
    rd.attach_answer_embeddings(docs)
    assert docs[1].answer_embeddings.shape == (2, rd.ANSWER_EMBEDDING_DIMENSIONS)

    docs[1].card = {**docs[1].card, "answers": ["可以加選幾學分？", "退選要怎麼辦？"]}
    rd.attach_answer_embeddings(docs)
    assert embedded[-1] == "退選要怎麼辦？" and len(embedded) == len(docs[0].card["answers"]) + 3  # 只算新的題目


def test_shortened_vectors_are_normalized():
    import numpy as np

    short = rd.shorten_vectors(np.array([[3.0, 4.0, 12.0]], dtype=np.float32), 2)
    assert np.allclose(short, [[0.6, 0.8]])


def test_concurrent_cache_writes_do_not_collide(tmp_path):
    # 評估跟網頁同時解析同一份文件時，兩邊會同時寫同一份快取
    import json
    from concurrent.futures import ThreadPoolExecutor

    target = tmp_path / "parsed" / "same.json"
    with ThreadPoolExecutor(max_workers=8) as pool:
        list(pool.map(lambda i: rd.write_json_cache(target, {"n": i}), range(40)))

    assert json.loads(target.read_text(encoding="utf-8"))["n"] in range(40)
    assert [p.name for p in target.parent.iterdir()] == ["same.json"]


def test_cache_replace_retries_while_the_target_is_locked(tmp_path, monkeypatch):
    # Windows 上目標檔正被別的行程讀取時，改名會「存取被拒」
    from pathlib import Path

    monkeypatch.setattr(rd.time, "sleep", lambda seconds: None)
    real_replace = Path.replace
    failures = iter([PermissionError("存取被拒"), PermissionError("存取被拒")])

    def flaky_replace(self, target):
        error = next(failures, None)
        if error:
            raise error
        return real_replace(self, target)

    monkeypatch.setattr(Path, "replace", flaky_replace)
    target = tmp_path / "card.json"
    rd.write_json_cache(target, {"ok": True})
    assert target.exists()


def test_cache_write_failure_does_not_break_the_answer(tmp_path, monkeypatch, caplog):
    from pathlib import Path

    monkeypatch.setattr(rd.time, "sleep", lambda seconds: None)

    def always_locked(self, target):
        raise PermissionError("存取被拒")

    monkeypatch.setattr(Path, "replace", always_locked)
    caplog.set_level("INFO", logger=rd.__name__)
    target = tmp_path / "card.json"
    rd.write_json_cache(target, {"ok": True})  # 不會丟例外
    assert not target.exists()
    assert list(tmp_path.iterdir()) == []  # 暫存檔也清掉
    assert "寫不進去" in caplog.text
